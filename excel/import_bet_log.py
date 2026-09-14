"""
Pulls new bet submissions from the site's Log Bet page (docs/log-bet.html)
and appends them into the Bet Log tab of MW_Handicapping_Tracker.xlsx.

WHY THIS EXISTS: docs/ is a plain static GitHub Pages site (no server -- see
README's "The website" section), so a browser form on it can never write
directly to a file on your machine. log-bet.html instead signs in with your
own Firebase Auth account and writes each submission straight to Firestore
(project mountain-dub-log-bet) -- that database is the durable landing spot.
Firestore's security rules only allow a signed-in write; nobody, signed in
or not, can read a bet back from the browser. This script is the other
half: it reads every "bets" document back out using a private
service-account key (never shipped to the browser -- see .env.example /
config.FIREBASE_SERVICE_ACCOUNT_PATH, which bypasses those rules for a
trusted server-side read) and appends each new row into the Bet Log tab's
columns A-H, exactly what you'd type in by hand:
    Date, Week, Matchup, Bet Type, Side, Line Taken, Odds (American), Stake (units)
Columns I (Closing Line) and K (Result) are deliberately left blank -- those
are only known later (once the market closes / the game is graded), same as
a hand-entered bet. Columns J/L/M are pre-built FORMULAS (CLV, Units
Won/Lost, Running Bankroll) already sitting on every template row and are
never touched here.

ONE-TIME SETUP:
  1. Firebase console -> gear icon (Project settings) -> Service accounts
     tab -> "Generate new private key". Save the downloaded JSON somewhere
     on your machine (NOT committed -- .gitignore already covers the
     recommended filename and Firebase's own default download name) and
     point FIREBASE_SERVICE_ACCOUNT_PATH at it in your .env (see
     .env.example).
  2. pip install -r requirements.txt (firebase-admin is already in it).

DEDUPING: Firestore mints a unique document id for every submission
(the addDoc() call in log-bet.html auto-generates it) -- this script keeps
a local record of which ids it's already imported
(excel/.bet_log_imported_ids.json, gitignored) so re-running it never
double-enters the same bet. Nothing is ever deleted from Firestore by this
script -- it stays a full history/backup of every submission, independent
of the tracker.

Usage:
    source .venv/bin/activate     (or the Windows equivalent)
    python excel/import_bet_log.py
"""
import json
import sys
import time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from config import FIREBASE_SERVICE_ACCOUNT_PATH  # noqa: E402

TRACKER_PATH = Path(__file__).resolve().parent / "MW_Handicapping_Tracker.xlsx"
IMPORTED_IDS_PATH = Path(__file__).resolve().parent / ".bet_log_imported_ids.json"
BET_LOG_ROWS = range(2, 43)  # matches build_tracker.py's Bet Log layout (row 1 = headers)

_firestore_client = None  # module-level cache so repeated calls in one run don't re-init the SDK


def _load_imported_ids():
    if not IMPORTED_IDS_PATH.exists():
        return set()
    return set(json.loads(IMPORTED_IDS_PATH.read_text()))


def _save_imported_ids(ids):
    IMPORTED_IDS_PATH.write_text(json.dumps(sorted(ids), indent=2))


def _get_firestore_client():
    global _firestore_client
    if _firestore_client is not None:
        return _firestore_client
    import firebase_admin
    from firebase_admin import credentials, firestore

    if not firebase_admin._apps:
        cred = credentials.Certificate(FIREBASE_SERVICE_ACCOUNT_PATH)
        firebase_admin.initialize_app(cred)
    _firestore_client = firestore.client()
    return _firestore_client


def fetch_submissions():
    db = _get_firestore_client()
    # Oldest first, same ordering the old Apps Script version used, so
    # rows land in the tracker in the order they were placed.
    docs = db.collection("bets").order_by("submitted_at").stream()
    rows = []
    for doc in docs:
        data = doc.to_dict()
        data["id"] = doc.id  # Firestore's own auto-id -- what we dedupe on
        rows.append(data)
    return rows


def first_empty_row(ws):
    for r in BET_LOG_ROWS:
        if ws[f"A{r}"].value in (None, ""):
            return r
    return None  # sheet is full -- see main()'s warning


def write_row(ws, row_num, sub):
    ws[f"A{row_num}"] = sub["date"]
    ws[f"B{row_num}"] = int(sub["week"])
    ws[f"C{row_num}"] = sub["matchup"]
    ws[f"D{row_num}"] = sub["bet_type"]
    ws[f"E{row_num}"] = sub["side"]
    ws[f"F{row_num}"] = float(sub["line"])
    ws[f"G{row_num}"] = int(sub["odds"])
    ws[f"H{row_num}"] = float(sub["stake"])
    # I (Closing Line) and K (Result) intentionally left blank. J/L/M are
    # pre-built formulas already sitting on this template row -- never
    # written here.


def main():
    if not FIREBASE_SERVICE_ACCOUNT_PATH:
        print("FIREBASE_SERVICE_ACCOUNT_PATH not set in .env -- see .env.example for setup.")
        return
    if not Path(FIREBASE_SERVICE_ACCOUNT_PATH).exists():
        print(f"Service account key not found at {FIREBASE_SERVICE_ACCOUNT_PATH}.")
        return
    if not TRACKER_PATH.exists():
        print(f"Tracker workbook not found at {TRACKER_PATH}.")
        return

    print("Fetching submissions from Firestore...")
    submissions = fetch_submissions()
    print(f"  {len(submissions)} total submission(s) in the 'bets' collection")

    imported_ids = _load_imported_ids()
    new_subs = [s for s in submissions if s["id"] not in imported_ids]
    if not new_subs:
        print("Nothing new to import.")
        return
    # already ordered oldest-first by the Firestore query in fetch_submissions()

    wb = openpyxl.load_workbook(TRACKER_PATH)  # formulas preserved as formulas, not evaluated
    ws = wb["Bet Log"]

    written = 0
    for sub in new_subs:
        row_num = first_empty_row(ws)
        if row_num is None:
            print(f"  Bet Log is full (rows {BET_LOG_ROWS.start}-{BET_LOG_ROWS.stop - 1}) -- "
                  f"stopping with {len(new_subs) - written} submission(s) still unimported. "
                  "Add more template rows (copy row 42's J/L/M formulas down) and re-run.")
            break
        write_row(ws, row_num, sub)
        imported_ids.add(sub["id"])
        print(f"  row {row_num}: {sub['side']} ({sub['bet_type']}), {sub['matchup']}")
        written += 1

    if written:
        wb.save(TRACKER_PATH)
        _save_imported_ids(imported_ids)
        print(f"\nImported {written} new bet(s) into Bet Log. Saved {TRACKER_PATH}. "
              "Open it in Excel -- formulas recalculate automatically on open.")
    else:
        print("Nothing written.")


if __name__ == "__main__":
    _script_start_time = time.time()
    print(f"[Started at {time.strftime('%Y-%m-%d %H:%M:%S')}]")
    main()

    _script_elapsed = time.time() - _script_start_time
    _mins, _secs = divmod(_script_elapsed, 60)
    print(f"\n[Finished in {int(_mins)}m {_secs:04.1f}s]" if _mins else f"\n[Finished in {_secs:.1f}s]")
