"""
Closing line value (CLV) tracking -- runs automatically as a run_pipeline.py step, so it
happens every time the GitHub workflow runs.

CLV asks: did you get a better NUMBER than the one the market closed at? Positive CLV on
a pick means the line moved toward your side after you "bought" it -- the most reliable
early sign of real skill, far less noisy than win/loss over a few hundred games.

Convention (the same one backtest.grade_spread_pick() uses): spreads are home-perspective,
negative = home favored; for a Home pick  CLV = entry - close, for an Away pick
CLV = close - entry. Positive is good.

WHAT THIS STEP DOES EACH RUN
  1. PICK LOG: appends Dub Gamma's current line + the consensus market spread for every
     FBS game in the upcoming week to data/clean/pick_log.csv (only when a game's numbers
     changed since the last logged row). This is "what the model showed when the workflow ran."
  2. CLV, two bases:
       "logged"  entry = the market spread the FIRST time the game appeared in the pick log;
                 close = the current/last consensus spread. This is the honest, live number and
                 only starts accumulating from the first run with this step.
       "open"    entry = the market's opening spread; the pick is Gamma's walk-forward lean
                 against that opening number. Computed for every graded game this season, so it
                 has data from day one -- a "would the model have beaten the opener" measure.
     Only games that have kicked off (or finished) count; their closing number is final.
  3. OUTPUTS: data/clean/clv_log.csv, docs/data/clv.json (the site's Tracking page card), and a
     "CLV" tab in the Excel tracker.
  4. BET LOG: fills the BLANK "Closing Line" cell of your own single-game Spread bets from the
     same consensus closing number once the game has kicked off. It never overwrites a closing
     line you typed in, and never touches any other column.

Read-only on the database. Everything is best-effort: a failure here prints a warning and the
pipeline carries on.
"""
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DB_PATH, CLEAN_DIR, ROOT  # noqa: E402
import backtest  # noqa: E402
import gamma_model  # noqa: E402
from predict_week import auto_detect_week  # noqa: E402
from teams import normalize_team_name, FBS_CONFERENCES  # noqa: E402

PICK_LOG = CLEAN_DIR / "pick_log.csv"
CLV_LOG = CLEAN_DIR / "clv_log.csv"
SITE_JSON = ROOT / "docs" / "data" / "clv.json"
TRACKER_PATH = ROOT / "excel" / "MW_Handicapping_Tracker.xlsx"
CLV_SHEET = "CLV"
LOG_COLS = ["run_ts", "season", "week", "game_id", "away_team", "home_team",
            "gamma_spread", "market_spread", "start_date"]


def _fbs_sql():
    return ",".join(f"'{c}'" for c in FBS_CONFERENCES)


def consensus_lines(con, season):
    return con.execute("""
        SELECT g.game_id, AVG(l.spread) AS close_spread, AVG(l.spread_open) AS open_spread
        FROM games g JOIN lines l ON l.game_id = g.game_id
        WHERE g.season = ? GROUP BY g.game_id
    """, [season]).fetchdf()


def week_games(con, season, week):
    return con.execute(f"""
        SELECT game_id, week, start_date, completed, home_team, away_team, home_points, away_points, neutral_site
        FROM games
        WHERE season = ? AND week = ? AND season_type = 'regular'
          AND home_conference IN ({_fbs_sql()}) AND away_conference IN ({_fbs_sql()})
    """, [season, week]).fetchdf()


def gamma_lines(con, season, week, games):
    ratings, _ = gamma_model.replay_ratings(con, season=season, through_week=week - 1)
    out = {}
    for r in games.itertuples():
        out[int(r.game_id)] = gamma_model.predict_spread_home(
            ratings, r.home_team, r.away_team, neutral_site=bool(r.neutral_site))
    return out


def clv_points(lean, entry, close):
    if lean == "Home":
        return entry - close
    if lean == "Away":
        return close - entry
    return None


def started(start_date, completed):
    if completed:
        return True
    try:
        return pd.Timestamp(start_date).tz_localize(None) <= pd.Timestamp.now("UTC").tz_localize(None)
    except Exception:
        return False


# ------------------------------------------------------------------ pick log
def log_current_picks(con):
    detected = auto_detect_week(con)
    if detected is None:
        print("  [CLV] no upcoming week found -- pick log unchanged.")
        return
    season, week = int(detected[0]), int(detected[1])
    games = week_games(con, season, week)
    if games.empty:
        return
    cons = consensus_lines(con, season).set_index("game_id")
    gl = gamma_lines(con, season, week, games)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    existing = pd.read_csv(PICK_LOG) if PICK_LOG.exists() else pd.DataFrame(columns=LOG_COLS)
    last = {}
    for r in existing.itertuples():
        last[int(r.game_id)] = (r.gamma_spread, r.market_spread)
    new_rows = []
    for r in games.itertuples():
        gid = int(r.game_id)
        mkt = cons["close_spread"].get(gid, np.nan)
        if pd.isna(mkt):
            continue
        g_spread, m_spread = round(float(gl[gid]), 2), round(float(mkt), 2)
        prev = last.get(gid)
        if prev is not None and abs(prev[0] - g_spread) < 0.01 and abs(prev[1] - m_spread) < 0.01:
            continue
        new_rows.append({"run_ts": ts, "season": season, "week": week, "game_id": gid,
                         "away_team": r.away_team, "home_team": r.home_team,
                         "gamma_spread": g_spread, "market_spread": m_spread,
                         "start_date": str(r.start_date)})
    if new_rows:
        pd.concat([existing, pd.DataFrame(new_rows)], ignore_index=True)[LOG_COLS].to_csv(PICK_LOG, index=False)
    print(f"  [CLV] pick log: week {week}, {len(new_rows)} new/changed row(s) ({len(existing) + len(new_rows)} total).")


# ------------------------------------------------------------------ CLV tables
def _grade_row(basis, season, week, game_id, away, home, gamma, entry, close, margin, completed):
    g = backtest.grade_spread_pick(float(gamma), float(entry), None,
                                   float(margin) if margin is not None and not pd.isna(margin) else 0.0,
                                   backtest.EDGE_THRESHOLD)
    return {"basis": basis, "season": season, "week": week, "game_id": game_id,
            "away_team": away, "home_team": home, "gamma_spread": round(float(gamma), 2),
            "entry_spread": round(float(entry), 2), "close_spread": round(float(close), 2),
            "lean": g["lean"], "edge": round(float(g["edge"]), 2), "is_bet": g["is_bet"],
            "clv_pts": None if g["lean"] == "Pick'em" else round(clv_points(g["lean"], entry, close), 2),
            "completed": bool(completed),
            "result": g["bet_result"] if completed and margin is not None and not pd.isna(margin) else None}


def build_logged_clv(con):
    if not PICK_LOG.exists():
        return []
    log = pd.read_csv(PICK_LOG)
    if log.empty:
        return []
    rows = []
    for season, sgrp in log.groupby("season"):
        cons = consensus_lines(con, int(season)).set_index("game_id")
        info = {r.game_id: r for r in con.execute(
            "SELECT game_id, start_date, completed, home_points, away_points FROM games WHERE season = ?",
            [int(season)]).fetchdf().itertuples()}
        first = sgrp.sort_values("run_ts").groupby("game_id").first()
        for gid, r in first.iterrows():
            gi = info.get(gid)
            if gi is None or gid not in cons.index or not started(gi.start_date, bool(gi.completed)):
                continue
            close = cons["close_spread"].get(gid)
            if pd.isna(close):
                continue
            margin = (gi.home_points - gi.away_points) if gi.completed and pd.notna(gi.home_points) else None
            rows.append(_grade_row("logged", int(season), int(r["week"]), int(gid), r["away_team"], r["home_team"],
                                   r["gamma_spread"], r["market_spread"], close, margin, bool(gi.completed)))
    return rows


def build_open_clv(con, season):
    cons = consensus_lines(con, season).set_index("game_id")
    wk_rows = con.execute("""
        SELECT DISTINCT week FROM games WHERE season = ? AND season_type = 'regular' AND completed = TRUE ORDER BY week
    """, [season]).fetchall()
    rows = []
    for (wk,) in wk_rows:
        if wk < 2:
            continue
        games = week_games(con, season, int(wk))
        games = games[games["completed"].astype(bool)]
        if games.empty:
            continue
        gl = gamma_lines(con, season, int(wk), games)
        for r in games.itertuples():
            gid = int(r.game_id)
            if gid not in cons.index:
                continue
            op, cl = cons["open_spread"].get(gid), cons["close_spread"].get(gid)
            if pd.isna(op) or pd.isna(cl):
                continue
            margin = r.home_points - r.away_points
            rows.append(_grade_row("open", season, int(wk), gid, r.away_team, r.home_team,
                                   gl[gid], op, cl, margin, True))
    return rows


def summarize(df):
    df = df.assign(clv_pts=pd.to_numeric(df["clv_pts"], errors="coerce"))
    out = {}
    for label, d in (("all", df), ("real_edge", df[df["is_bet"] == True])):  # noqa: E712
        d = d[d["clv_pts"].notna()]
        n = len(d)
        out[label] = {
            "games": n,
            "mean_clv": round(float(d["clv_pts"].mean()), 3) if n else None,
            "pct_positive": round(float((d["clv_pts"] > 0).mean()), 3) if n else None,
            "pct_negative": round(float((d["clv_pts"] < 0).mean()), 3) if n else None,
            "pct_unchanged": round(float((d["clv_pts"] == 0).mean()), 3) if n else None,
        }
    return out


# ------------------------------------------------------------------ outputs
def write_site_json(clv_df, summaries):
    SITE_JSON.parent.mkdir(parents=True, exist_ok=True)
    recent = []
    if not clv_df.empty:
        rec = clv_df[clv_df["basis"] == "logged"].sort_values(["week", "game_id"], ascending=False).head(25)
        recent = json.loads(rec.where(rec.notna(), None).to_json(orient="records"))
    payload = {"generated": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
               "bases": summaries, "recent_logged": recent}
    SITE_JSON.write_text(json.dumps(payload, indent=1, default=str), encoding="utf-8")


def fill_bet_log(con, wb):
    from openpyxl.utils import get_column_letter  # noqa: F401
    if "Bet Log" not in wb.sheetnames:
        return 0
    ws = wb["Bet Log"]
    header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    need = ["Matchup", "Bet Type", "Side", "Line Taken", "Closing Line", "Week"]
    if any(n not in header for n in need):
        print("  [CLV] Bet Log columns not as expected -- closing lines not auto-filled.")
        return 0
    season = con.execute("SELECT MAX(season) FROM games WHERE completed = TRUE").fetchone()[0]
    cons = consensus_lines(con, season).set_index("game_id")
    gm = {}
    for r in con.execute("SELECT game_id, home_team, away_team, start_date, completed FROM games WHERE season = ?",
                         [season]).fetchall():
        gm[(normalize_team_name(r[2]).lower(), normalize_team_name(r[1]).lower())] = r
    filled = 0
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, header["Bet Type"]).value or "").strip().lower() != "spread":
            continue
        if ws.cell(row, header["Closing Line"]).value not in (None, ""):
            continue
        matchup = str(ws.cell(row, header["Matchup"]).value or "")
        if "@" not in matchup or "," in matchup:
            continue
        away, home = [p.strip() for p in matchup.split("@", 1)]
        g = gm.get((normalize_team_name(away).lower(), normalize_team_name(home).lower()))
        if g is None or g[0] not in cons.index or not started(g[3], bool(g[4])):
            continue
        close_home = cons["close_spread"].get(g[0])
        if pd.isna(close_home):
            continue
        side = normalize_team_name(str(ws.cell(row, header["Side"]).value or "")).lower()
        if side == normalize_team_name(home).lower():
            close_side = close_home
        elif side == normalize_team_name(away).lower():
            close_side = -close_home
        else:
            continue
        ws.cell(row, header["Closing Line"]).value = round(close_side * 2) / 2.0
        filled += 1
    return filled


def write_clv_tab(wb, clv_df, summaries):
    from openpyxl.styles import Font, PatternFill, Alignment
    if CLV_SHEET in wb.sheetnames:
        del wb[CLV_SHEET]
    ws = wb.create_sheet(CLV_SHEET)
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws["A1"] = "Closing Line Value -- Dub Gamma picks"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("CLV (pts) = entry line - closing line for a Home pick, closing - entry for an Away pick. "
                "Positive = you got a better number than the close. 'logged' = the market line when the workflow first "
                "saw the game; 'open' = the market's opening line. Refreshed on every pipeline run.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 48
    r = 4
    for basis, title in (("logged", "Logged entry (live)"), ("open", "Opening line")):
        s = summaries.get(basis)
        ws.cell(r, 1, title).font = Font(bold=True)
        r += 1
        for c, h in enumerate(["Slice", "Games", "Avg CLV (pts)", "% Positive", "% Negative", "% Unchanged"], 1):
            cell = ws.cell(r, c, h)
            cell.fill, cell.font = hdr_fill, hdr_font
        r += 1
        for label, name in (("all", "All picks"), ("real_edge", "Real edge (>= 2 pts)")):
            d = (s or {}).get(label) or {}
            ws.cell(r, 1, name)
            ws.cell(r, 2, d.get("games", 0))
            ws.cell(r, 3, d.get("mean_clv"))
            for c, k in ((4, "pct_positive"), (5, "pct_negative"), (6, "pct_unchanged")):
                v = d.get(k)
                cell = ws.cell(r, c, v)
                cell.number_format = "0.0%"
            r += 1
        r += 1
    r += 1
    cols = ["basis", "season", "week", "away_team", "home_team", "lean", "gamma_spread", "entry_spread",
            "close_spread", "edge", "is_bet", "clv_pts", "result"]
    heads = ["Basis", "Season", "Week", "Away", "Home", "Gamma Lean", "Gamma Line (Home)", "Entry Line (Home)",
             "Close Line (Home)", "Edge (pts)", "Real Edge", "CLV (pts)", "Result vs Entry"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(r, c, h)
        cell.fill, cell.font = hdr_fill, hdr_font
    first_data = r + 1
    if not clv_df.empty:
        d = clv_df.sort_values(["basis", "week", "game_id"], ascending=[True, False, True])
        for i, rec in enumerate(d[cols].itertuples(index=False), start=first_data):
            for c, v in enumerate(rec, 1):
                if isinstance(v, (np.bool_, bool)):
                    v = "Yes" if v else ""
                elif isinstance(v, float) and np.isnan(v):
                    v = None
                elif isinstance(v, np.generic):
                    v = v.item()
                ws.cell(i, c, v)
    for col, w in zip("ABCDEFGHIJKLM", (10, 8, 7, 24, 24, 12, 16, 16, 16, 11, 10, 10, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(first_data, 1)


def _run():
    start = time.time()
    CLEAN_DIR.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(str(DB_PATH))
    try:
        log_current_picks(con)
    except Exception as e:  # never break the pipeline over this
        print(f"  [CLV] pick log skipped: {type(e).__name__}: {e}")
    rows = []
    try:
        rows += build_logged_clv(con)
    except Exception as e:
        print(f"  [CLV] logged CLV skipped: {type(e).__name__}: {e}")
    try:
        season = con.execute("SELECT MAX(season) FROM games WHERE completed = TRUE").fetchone()[0]
        rows += build_open_clv(con, int(season))
    except Exception as e:
        print(f"  [CLV] open-line CLV skipped: {type(e).__name__}: {e}")
    clv_df = pd.DataFrame(rows)
    if not clv_df.empty:
        clv_df.to_csv(CLV_LOG, index=False)
    summaries = {b: summarize(clv_df[clv_df["basis"] == b]) for b in ("logged", "open")} if not clv_df.empty else {}
    write_site_json(clv_df, summaries)
    for b, s in summaries.items():
        a = s["all"]
        print(f"  [CLV] {b}: {a['games']} games, avg CLV {a['mean_clv']} pts, "
              f"{(a['pct_positive'] or 0):.0%} positive")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(TRACKER_PATH)
        filled = fill_bet_log(con, wb)
        write_clv_tab(wb, clv_df, summaries)
        wb.save(TRACKER_PATH)
        print(f"  [CLV] Excel: CLV tab refreshed; {filled} blank Bet Log closing line(s) filled.")
    except PermissionError:
        print("  [CLV] Excel tracker is open/locked -- CLV tab and Bet Log fill skipped this run.")
    except Exception as e:
        print(f"  [CLV] Excel update skipped: {type(e).__name__}: {e}")
    con.close()
    print(f"  [CLV] done in {time.time() - start:.1f}s")


def main():
    try:
        _run()
    except Exception as e:  # CLV is a bonus step -- never let it stop the pipeline
        print(f"  [CLV] step failed and was skipped: {type(e).__name__}: {e}")


if __name__ == "__main__":
    main()
